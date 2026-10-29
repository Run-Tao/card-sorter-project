# -*- coding: utf-8 -*-
"""生成 pandoc 参考样式文档（中文报告用）。"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "tools", "reference.docx")

doc = Document()
for s in doc.sections:
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.top_margin = s.bottom_margin = Cm(2.2)
    s.left_margin = s.right_margin = Cm(2.2)


def set_style(name, cn="宋体", ascii_f="Times New Roman", size=10.5,
              bold=None, color=None, before=None, after=None,
              line=None, keep_next=None, align=None):
    try:
        st = doc.styles[name]
    except KeyError:
        return
    st.font.name = ascii_f
    st.font.size = Pt(size)
    if bold is not None:
        st.font.bold = bold
    if color is not None:
        st.font.color.rgb = color
    rpr = st.element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    rf.set(qn("w:eastAsia"), cn)
    pf = st.paragraph_format
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line
    if keep_next is not None:
        pf.keep_with_next = keep_next
    if align is not None:
        pf.alignment = align


BLUE = RGBColor(0x1F, 0x38, 0x64)
MID = RGBColor(0x2F, 0x5D, 0x8A)

set_style("Normal", "宋体", "Times New Roman", 10.5, line=1.35, after=3)
set_style("Body Text", "宋体", "Times New Roman", 10.5, line=1.35, after=3)
set_style("First Paragraph", "宋体", "Times New Roman", 10.5, line=1.35, after=3)
set_style("Compact", "宋体", "Times New Roman", 9, line=1.15, after=1)
set_style("Title", "微软雅黑", "Arial", 22, bold=True, color=BLUE, after=6,
          align=WD_ALIGN_PARAGRAPH.CENTER)
set_style("Subtitle", "微软雅黑", "Arial", 13, bold=False, color=MID, after=10,
          align=WD_ALIGN_PARAGRAPH.CENTER)
set_style("Date", "宋体", "Times New Roman", 10.5, color=RGBColor(0x59, 0x59, 0x59),
          align=WD_ALIGN_PARAGRAPH.CENTER)
set_style("Author", "宋体", "Times New Roman", 10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
set_style("Heading 1", "微软雅黑", "Arial", 16, bold=True, color=BLUE,
          before=18, after=8, keep_next=True)
set_style("Heading 2", "微软雅黑", "Arial", 13.5, bold=True, color=MID,
          before=12, after=5, keep_next=True)
set_style("Heading 3", "微软雅黑", "Arial", 11.5, bold=True, color=RGBColor(0x33, 0x33, 0x33),
          before=8, after=3, keep_next=True)
set_style("Heading 4", "微软雅黑", "Arial", 10.5, bold=True, color=RGBColor(0x33, 0x33, 0x33),
          before=6, after=2, keep_next=True)
set_style("Table", "宋体", "Times New Roman", 9, line=1.1, after=0)
set_style("Table Caption", "宋体", "Times New Roman", 9, color=RGBColor(0x55, 0x55, 0x55),
          align=WD_ALIGN_PARAGRAPH.CENTER)
set_style("Image Caption", "宋体", "Times New Roman", 9, color=RGBColor(0x55, 0x55, 0x55),
          align=WD_ALIGN_PARAGRAPH.CENTER)
set_style("Source Code", "Consolas", "Consolas", 8.5, line=1.0)
set_style("Verbatim Char", "Consolas", "Consolas", 8.5)

# 页脚：文档名 + 页码
footer = doc.sections[0].footer
p = footer.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("无标记扑克牌自动分拣机 · 工程实施方案 v1.0    第 ")
run.font.size = Pt(8.5)
run.font.name = "宋体"
run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

fld1 = OxmlElement("w:fldChar")
fld1.set(qn("w:fldCharType"), "begin")
instr = OxmlElement("w:instrText")
instr.set(qn("xml:space"), "preserve")
instr.text = "PAGE"
fld2 = OxmlElement("w:fldChar")
fld2.set(qn("w:fldCharType"), "end")
r = p.add_run()
r._r.append(fld1)
r._r.append(instr)
r._r.append(fld2)
r.font.size = Pt(8.5)

run2 = p.add_run(" 页")
run2.font.size = Pt(8.5)
run2.font.name = "宋体"
run2._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

doc.save(OUT)
print("saved", OUT)

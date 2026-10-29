# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Emu

p = r"report\无标记扑克牌自动分拣机_工程实施方案.docx"
d = Document(p)
print("paragraphs", len(d.paragraphs), "tables", len(d.tables), "images", len(d.inline_shapes))
sec = d.sections[0]
print("page width cm", round(sec.page_width.cm, 2),
      "usable", round(Emu(sec.page_width - sec.left_margin - sec.right_margin).cm, 2))
for i, sh in enumerate(d.inline_shapes):
    print(" image", i, "w_cm", round(sh.width.cm, 2), "h_cm", round(sh.height.cm, 2))
print("--- first 30 paragraphs ---")
for para in d.paragraphs[:30]:
    t = para.text.strip()
    if t:
        print(f"[{para.style.name}] {t[:100]}")
print("--- table dimensions ---")
for i, tb in enumerate(d.tables):
    print(i, len(tb.rows), "x", len(tb.columns))
